"""Agent F: wiring of the PDF / PowerPoint exports: API routes (media types, file names, per-language cache with the
report's invalidation stamp), CLI `report --format`, the export zip, the e-mail attachment, the suspicious-rows
section in all three outputs, and the generation time on the large run (skipped when workspace/te_2m is absent)."""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("reportlab")
pytest.importorskip("pptx")
pypdf = pytest.importorskip("pypdf")

from pptx import Presentation  # noqa: E402

from tests.fixtures.fake_workspace_f import build_fake_workspace  # noqa: E402
from tpm.contracts import HumanDecision  # noqa: E402
from tpm.report import Translator, ensure_export, generate_report, report_path  # noqa: E402
from tpm.report.export_common import MEDIA_TYPES, download_name, export_context  # noqa: E402
from tpm.report.pdf import generate_pdf  # noqa: E402
from tpm.report.pptx_export import generate_pptx  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
RUN = "run_export_f"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _pdf_text(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    return re.sub(r"\s+", " ", "\n".join(p.extract_text() or "" for p in reader.pages))


def _deck_text(prs) -> str:
    out = []
    for sl in prs.slides:
        for s in sl.shapes:
            if s.has_text_frame:
                out.append(s.text_frame.text)
            if getattr(s, "has_table", False) and s.has_table:
                out += [c.text_frame.text for r in s.table.rows for c in r.cells]
    return re.sub(r"\s+", " ", " ".join(out))


@pytest.fixture
def fake_ws(tmp_path, monkeypatch):
    monkeypatch.delenv("TPM_REPORT_LLM", raising=False)
    ws = build_fake_workspace(tmp_path / "workspace")
    yield ws
    ws.close()


# ----------------------------------------------------------------------------- API
@pytest.fixture(scope="module")
def client(tmp_path_factory):
    from fastapi.testclient import TestClient

    from tpm.api.server import create_app
    from tpm.config import load_settings

    base = tmp_path_factory.mktemp("f_export_api")
    settings_path = base / "settings.yaml"
    shutil.copy(ROOT / "config" / "settings.yaml", settings_path)
    s = load_settings(settings_path)
    s.workspace_dir = str(base / "workspace")
    ws = build_fake_workspace(base / "workspace", run_id=RUN, settings=s)
    ws.close()
    app = create_app(settings_path=settings_path, workspace_dir=base / "workspace")
    with TestClient(app) as c:
        yield c


def test_routes_media_types_names_and_cache(client):
    r = client.get(f"/api/runs/{RUN}/report.pdf", params={"lang": "en"})
    assert r.status_code == 200, r.text[:300]
    assert r.headers["content-type"].startswith("application/pdf") and r.content[:5] == b"%PDF-"
    assert f'filename="tpm_{RUN}_en.pdf"' in r.headers["content-disposition"]
    assert r.headers["x-tpm-export-regenerated"] == "1" and r.headers["cache-control"] == "no-store"
    assert Translator("en")("section_8") in _pdf_text(r.content)
    again = client.get(f"/api/runs/{RUN}/report.pdf", params={"lang": "en"})
    assert again.status_code == 200 and again.headers["x-tpm-export-regenerated"] == "0" and again.content == r.content  # served from the cache

    p = client.get(f"/api/runs/{RUN}/report.pptx", params={"lang": "fi"})
    assert p.status_code == 200 and p.headers["content-type"].startswith(PPTX_TYPE) and p.content[:2] == b"PK"
    assert f'filename="tpm_{RUN}_fi.pptx"' in p.headers["content-disposition"]
    prs = Presentation(io.BytesIO(p.content))
    assert 10 <= len(prs.slides) <= 16 and Translator("fi")("px_trust") in _deck_text(prs)
    assert client.get(f"/api/runs/{RUN}/report.pdf", params={"lang": "fi"}).headers["x-tpm-export-regenerated"] == "1"  # cached per language

    # a human decision changes the report stamp: the next download is rebuilt and shows the decision
    d = client.post(f"/api/runs/{RUN}/decisions", json={"actor_name": "exporttester", "role": "reviewer", "action": "question", "object_type": "flag", "object_id": "FLAG-000002", "note": "checked for the export test"})
    assert d.status_code in (200, 201), d.text[:300]
    fresh = client.get(f"/api/runs/{RUN}/report.pdf", params={"lang": "en"})
    assert fresh.headers["x-tpm-export-regenerated"] == "1" and "exporttester" in _pdf_text(fresh.content)
    forced = client.get(f"/api/runs/{RUN}/report.pdf", params={"lang": "en", "refresh": 1})
    assert forced.headers["x-tpm-export-regenerated"] == "1"

    assert client.get(f"/api/runs/{RUN}/report.pdf", params={"lang": "xx"}).status_code == 200  # unknown language -> default
    assert client.get("/api/runs/no_such_run/report.pdf").status_code == 404
    assert client.get("/api/runs/no_such_run/report.pptx").status_code == 404
    html = client.get(f"/api/runs/{RUN}/report", params={"lang": "en"})  # the HTML route is untouched by the new ones
    assert html.status_code == 200 and "<html" in html.text.lower()


def test_report_view_has_the_download_buttons():
    js = (ROOT / "tpm" / "api" / "static" / "js" / "views" / "report.js").read_text(encoding="utf-8")
    assert "/report.${kind}" in js and "rep.downloadPdf" in js and "rep.downloadPptx" in js
    assert "rep.exportGenerating" in js and "rep.exportFailed" in js  # generating state and error feedback
    for lang in ("en", "fi", "sv"):
        d = json.loads((ROOT / "tpm" / "api" / "static" / "i18n" / f"{lang}.json").read_text(encoding="utf-8"))
        for key in ("rep.downloadPdf", "rep.downloadPptx", "rep.exportGenerating", "rep.exportWorking", "rep.exportReady", "rep.exportFailed", "rep.exportHelp"):
            assert d.get(key), f"{lang}: {key}"


# ----------------------------------------------------------------------------- cache
def test_ensure_export_cache_follows_the_report_stamp(fake_ws, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("an export must never call the language model")

    monkeypatch.setattr("tpm.llm.complete", boom)
    first = ensure_export(fake_ws, fake_ws.settings, "en", "pdf")
    assert first["regenerated"] and first["media_type"] == MEDIA_TYPES["pdf"] and first["filename"] == download_name(fake_ws.run_id, "en", "pdf")
    assert Path(first["path"]).name == "report_en.pdf" and first["bytes"] > 40_000
    assert not ensure_export(fake_ws, fake_ws.settings, "en", "pdf")["regenerated"]
    assert ensure_export(fake_ws, fake_ws.settings, "en", "pptx")["regenerated"]  # formats are cached separately
    assert not ensure_export(fake_ws, fake_ws.settings, "en", "pptx")["regenerated"]
    assert ensure_export(fake_ws, fake_ws.settings, "en", "pdf", force=True)["regenerated"]
    # an artifact changes -> stale
    flags = [f.model_dump() for f in fake_ws.flags()]
    flags[0]["statement"] = "A statement that changed after the first export."
    time.sleep(0.02)
    fake_ws.rewrite_jsonl("flags", flags)
    again = ensure_export(fake_ws, fake_ws.settings, "en", "pdf")
    assert again["regenerated"] and "A statement that changed after the first export." in _pdf_text(Path(again["path"]).read_bytes())
    # a human decision -> stale
    from tpm.pipeline import apply_decision

    apply_decision(fake_ws, fake_ws.settings, HumanDecision(actor_name="outi", role="reviewer", action="accept", object_type="diagnosis", object_id="DIAG-000002"))
    assert ensure_export(fake_ws, fake_ws.settings, "en", "pdf")["regenerated"]
    with pytest.raises(ValueError):
        ensure_export(fake_ws, fake_ws.settings, "en", "docx")
    assert fake_ws.log.verify_chain()["ok"] and any(e.action == "report_export" for e in fake_ws.log.entries(action="report_export"))


# ----------------------------------------------------------------------------- suspicious rows (lead's artifact)
SUSPICIOUS = {
    "headline": "31 rows stand out as suspicious; 27 of them are isolated readings.",
    "wording": "glitch or manipulation: the data alone can't tell",
    "regime": {"point_dominated": True, "n_point_stretches": 27, "n_sustained_stretches": 4, "share_points": 0.87},
    "n_rows": 31, "n_listed": 25, "cap": 25,
    "rows": [{"row": 100 + 7 * i, "row_end": 100 + 7 * i + (2 if i % 5 == 0 else 0), "group_id": str(1 + i % 6), "batch_id": f"B{i // 5:04d}", "signals": [{"signal": f"S{1 + i % 9:02d}", "deviation": 9.3 - 0.2 * i, "direction": "up" if i % 2 else "down", "explanation": f"S{1 + i % 9:02d} jumped for a single reading and came straight back."}], "sources": ["detector", "local_spike"] if i % 3 else ["range_check"], "strength": round(0.99 - 0.02 * i, 2), "flag_ids": ["FLAG-000090"] if i == 0 else [], "check_ids": ["CHK-000004"] if i % 3 == 0 else [], "evidence_ids": ["EV-000003"], "statement": f"Row {100 + 7 * i}: a reading far outside what its neighbours show."} for i in range(25)],
}


def _add_suspicious(ws):
    ws.write_json("suspicious_rows.json", SUSPICIOUS)
    flags = [f.model_dump() for f in ws.flags()]
    point = dict(flags[0])
    point.update({"id": "FLAG-000090", "kind": "point", "row_start": 100, "row_end": 100, "pattern_id": "PATTERN-A", "severity": 0.99, "statement": "Isolated reading at row 100: S01 far outside its neighbours."})
    ws.rewrite_jsonl("flags", flags + [point])
    diags = [d.model_dump() for d in ws.diagnoses()]
    pd_ = dict(diags[0])
    pd_.update({"id": "DIAG-000090", "flag_ids": ["FLAG-000090"], "fault_type": "isolated reading on S01", "steps": ["When it started: Onset at row 100.", "What changed first: S01 jumped.", "How it propagated: nothing followed."]})
    ws.rewrite_jsonl("diagnoses", diags + [pd_])


def test_suspicious_rows_section_in_html_pdf_and_deck(fake_ws):
    t = Translator("en")
    # absent -> omitted everywhere
    html = generate_report(fake_ws, fake_ws.settings, "en", use_llm=False).read_text(encoding="utf-8")
    assert 'id="suspicious"' not in html and t("section_suspicious") not in html
    assert t("section_suspicious") not in _pdf_text(generate_pdf(fake_ws, fake_ws.settings, "en").read_bytes())
    n_before = len(Presentation(str(generate_pptx(fake_ws, fake_ws.settings, "en"))).slides)

    _add_suspicious(fake_ws)
    ctx = export_context(fake_ws, fake_ws.settings, "en")
    assert ctx["suspicious"]["shown"] == 25 and ctx["suspicious"]["n_point_flags"] == 1
    assert all(f["kind"] != "point" for f in ctx["detect"]["flags"]) and ctx["detect"]["n_flags"] == len(fake_ws.flags()) - 1
    card = next(d for d in ctx["diagnoses"] if d["id"] == "DIAG-000090")
    assert card["point_only"] and not card["propagation"] and card["pattern_id"] is None and card["steps"] == ["What changed first: S01 jumped."]

    html = generate_report(fake_ws, fake_ws.settings, "en", use_llm=False).read_text(encoding="utf-8")
    sect = html.split('id="suspicious"', 1)[1].split('id="section-1"', 1)[0]
    assert html.count('<section id="suspicious"') == 1  # ONE headline list
    assert SUSPICIOUS["headline"] in sect and "glitch or manipulation: the data alone can&#39;t tell" in sect and 'class="wording"' in sect
    assert sect.count("a reading far outside what its neighbours show") == 25 and "FLAG-000090" in sect and "local spike" in sect and "range check" in sect
    flags_table = html.split('<table class="fixed flags">', 1)[1].split("</table>", 1)[0]
    assert "FLAG-000090" not in flags_table  # a point flag is not a sustained event

    pdf = _pdf_text(generate_pdf(fake_ws, fake_ws.settings, "en").read_bytes())
    assert t("section_suspicious") in pdf and SUSPICIOUS["wording"] in pdf and SUSPICIOUS["headline"] in pdf
    assert pdf.count("a reading far outside what its neighbours show") == 25 and t("susp_point_flags", n=1) in pdf

    prs = Presentation(str(generate_pptx(fake_ws, fake_ws.settings, "en")))
    assert len(prs.slides) == n_before + 1 <= 16
    slide = next(sl for sl in prs.slides if any(s.name.startswith("tpm-wording") for s in sl.shapes))
    assert next(s for s in slide.shapes if s.name.startswith("tpm-wording")).text_frame.text == SUSPICIOUS["wording"]
    rows = [r for s in slide.shapes if getattr(s, "has_table", False) and s.has_table for r in list(s.table.rows)[1:]]
    assert len(rows) == 25 and rows[0].cells[0].text_frame.text == "100–102"
    assert "FLAG-000090" in slide.notes_slide.notes_text_frame.text

    fi = _pdf_text(generate_pdf(fake_ws, fake_ws.settings, "fi").read_bytes())
    assert Translator("fi")("section_suspicious") in fi and Translator("fi")("susp_source_range_check") in fi


# ----------------------------------------------------------------------------- export zip, e-mail, CLI
def test_export_zip_contains_pdf_and_deck(fake_ws, tmp_path):
    from tpm.log.exports import export_run, list_export

    zp = export_run(fake_ws, tmp_path / "out", languages=["fi"])
    names = list_export(zp)
    for must in ("report_fi.html", "report_fi.pdf", "report_fi.pptx", "decision_log.jsonl", "manifest.json"):
        assert must in names, must
    assert "EXPORT_ERROR.txt" not in names and "dataset.parquet" not in names
    with zipfile.ZipFile(zp) as z:
        assert z.read("report_fi.pdf")[:5] == b"%PDF-" and z.read("report_fi.pptx")[:2] == b"PK"
    without = list_export(export_run(fake_ws, tmp_path / "out2", languages=["fi"], make_exports=False))
    assert "report_fi.html" in without and not [n for n in without if n.endswith((".pdf", ".pptx"))]


def test_email_can_attach_the_pdf(fake_ws, monkeypatch):
    sent = {}

    class FakeSMTP:
        def __init__(self, host, port, timeout=None):
            pass

        def starttls(self):
            pass

        def login(self, u, p):
            pass

        def send_message(self, msg):
            sent["msg"] = msg

        def quit(self):
            pass

    import smtplib

    from tpm.report import email_report

    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    monkeypatch.setenv("TPM_SMTP_HOST", "smtp.example.org")
    monkeypatch.setenv("TPM_SMTP_PORT", "587")
    res = email_report(fake_ws, fake_ws.settings, "x@y.z", "sv", attach_pdf=True)
    atts = list(sent["msg"].iter_attachments())
    assert [a.get_filename() for a in atts] == ["report_sv.html", f"tpm_{fake_ws.run_id}_sv.pdf"] and res["attachments"] == ["report_sv.html", "report_sv.pdf"]
    assert atts[1].get_content_type() == "application/pdf" and atts[1].get_payload(decode=True)[:5] == b"%PDF-"
    email_report(fake_ws, fake_ws.settings, "x@y.z", "sv")  # the default stays HTML only
    assert len(list(sent["msg"].iter_attachments())) == 1


def _fake_smtp(sent: dict, fail: Exception | None = None):
    class FakeSMTP:
        def __init__(self, host, port, timeout=None):
            sent.setdefault("connections", []).append((type(self).__name__, host, port))

        def starttls(self):
            sent["starttls"] = True

        def login(self, u, p):
            sent["login"] = (u, p)

        def send_message(self, msg):
            if fail is not None:
                raise fail
            sent["msg"] = msg

        def quit(self):
            pass

    return FakeSMTP


def test_email_attaches_pdf_and_deck_over_implicit_tls(fake_ws, monkeypatch):
    """The Report view's default: HTML + PDF + deck in one message, through an SMTPS port (Resend: 465 or 2465)."""
    import smtplib

    from tpm.report import email_report

    sent: dict = {}
    fake = _fake_smtp(sent)
    monkeypatch.setattr(smtplib, "SMTP_SSL", type("SMTP_SSL", (fake,), {}))
    monkeypatch.setattr(smtplib, "SMTP", type("SMTP", (fake,), {}))
    monkeypatch.setenv("TPM_SMTP_HOST", "smtp.resend.com")
    monkeypatch.setenv("TPM_SMTP_PORT", "2465")
    monkeypatch.setenv("TPM_SMTP_USER", "resend")
    monkeypatch.setenv("TPM_SMTP_PASSWORD", "re_test_not_a_key")
    monkeypatch.setenv("TPM_SMTP_FROM", "onboarding@resend.dev")
    res = email_report(fake_ws, fake_ws.settings, "me@example.com", "en", attach_pdf=True, attach_pptx=True)
    assert sent["connections"] == [("SMTP_SSL", "smtp.resend.com", 2465)] and "starttls" not in sent
    assert sent["login"] == ("resend", "re_test_not_a_key") and sent["msg"]["From"] == "onboarding@resend.dev"
    atts = list(sent["msg"].iter_attachments())
    assert [a.get_filename() for a in atts] == ["report_en.html", f"tpm_{fake_ws.run_id}_en.pdf", f"tpm_{fake_ws.run_id}_en.pptx"]
    assert [a.get_content_type() for a in atts] == ["text/html", "application/pdf", PPTX_TYPE]
    assert atts[1].get_payload(decode=True)[:5] == b"%PDF-" and atts[2].get_payload(decode=True)[:2] == b"PK"
    assert res["attachments"] == ["report_en.html", "report_en.pdf", "report_en.pptx"]
    log = fake_ws.log.entries(action="email")
    assert log and log[-1].payload["attachments"] == res["attachments"] and "re_test_not_a_key" not in json.dumps(log[-1].payload)


def test_email_failure_names_the_mail_servers_reason(fake_ws, monkeypatch):
    import smtplib

    from tpm.report import email_report

    refused = smtplib.SMTPDataError(550, b"You can only send testing emails to your own email address (owner@example.com).")
    monkeypatch.setattr(smtplib, "SMTP", _fake_smtp({}, fail=refused))
    monkeypatch.setenv("TPM_SMTP_HOST", "smtp.resend.com")
    monkeypatch.setenv("TPM_SMTP_PORT", "587")
    with pytest.raises(RuntimeError) as ei:
        email_report(fake_ws, fake_ws.settings, "someone.else@example.com", "en")
    assert "550 You can only send testing emails to your own email address" in str(ei.value) and "b'" not in str(ei.value)

    def unreachable(*a, **k):
        raise TimeoutError("timed out")

    monkeypatch.setattr(smtplib, "SMTP_SSL", unreachable)
    monkeypatch.setenv("TPM_SMTP_PORT", "465")
    with pytest.raises(RuntimeError) as ei:
        email_report(fake_ws, fake_ws.settings, "me@example.com", "en")
    assert "smtp.resend.com:465" in str(ei.value) and "2587" in str(ei.value)


def test_cli_report_formats(fake_ws, tmp_path):
    env = dict(os.environ)
    env.update({"TPM_WORKSPACE": str(fake_ws.dir.parent), "PYTHONIOENCODING": "utf-8", "TPM_REPORT_LLM": "0"})
    env.pop("TPM_PROFILE", None)
    out_dir = tmp_path / "cli_out"
    r = subprocess.run([sys.executable, "-m", "tpm", "report", fake_ws.run_id, "--format", "all", "--lang", "en", "--out", str(out_dir), "--no-llm"], cwd=str(ROOT), env=env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=240)
    assert r.returncode == 0, r.stderr[-800:]
    for ext, magic in (("html", b"<!DOCTYPE"), ("pdf", b"%PDF-"), ("pptx", b"PK")):
        p = out_dir / f"tpm_{fake_ws.run_id}_en.{ext}"
        assert p.exists() and p.read_bytes()[: len(magic)] == magic, ext
    assert "PDF written" in r.stdout and "PowerPoint written" in r.stdout
    single = tmp_path / "one.pdf"
    r = subprocess.run([sys.executable, "-m", "tpm", "report", fake_ws.run_id, "--format", "pdf", "--lang", "fi", "--out", str(single)], cwd=str(ROOT), env=env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=240)
    assert r.returncode == 0 and single.read_bytes()[:5] == b"%PDF-", r.stderr[-800:]
    from tpm.cli import build_parser

    ns = build_parser().parse_args(["report", "x"])
    assert ns.format == "html" and ns.pdf_engine == "native"  # the default behaviour of `tpm report` is unchanged
    assert build_parser().parse_args(["email", "x", "--to", "a@b.c", "--pdf"]).pdf is True


# ----------------------------------------------------------------------------- large run
@pytest.mark.skipif(not (ROOT / "workspace" / "te_2m" / "status.json").exists(), reason="workspace/te_2m is not present")
def test_both_exports_finish_in_time_on_the_large_run(tmp_path):
    from tpm.config import load_settings
    from tpm.workspace import Workspace

    s = load_settings()
    s.workspace_dir = str(ROOT / "workspace")
    ws = Workspace.open("te_2m", s)
    try:
        t0 = time.time()
        pdf = generate_pdf(ws, s, "en", out_path=tmp_path / "te_2m.pdf", log=False)  # nothing is written into the run
        deck = generate_pptx(ws, s, "en", out_path=tmp_path / "te_2m.pptx", log=False)
        seconds = time.time() - t0
    finally:
        ws.close()
    assert seconds < 30, f"PDF + PowerPoint took {seconds:.1f} s on te_2m"
    reader = pypdf.PdfReader(str(pdf))
    assert 20 <= len(reader.pages) <= 120 and pdf.stat().st_size < 3_000_000
    text = re.sub(r"\s+", " ", "\n".join(p.extract_text() or "" for p in reader.pages[:70]))
    t = Translator("en")
    assert re.search(r"Showing the \d+ most severe of [\d,]+ flags", text), "row caps state the total"
    prs = Presentation(str(deck))
    assert 10 <= len(prs.slides) <= 16 and deck.stat().st_size < 3_000_000
    assert t("px_evaluation") in _deck_text(prs)
