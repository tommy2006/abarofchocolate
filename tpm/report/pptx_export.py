"""PowerPoint export of the run report (agent F): a 16:9 deck a team can present as it is.

    generate_pptx(ws, settings, lang="en", out_path=None) -> Path
    render_pptx(context) -> bytes

About 10 to 16 slides built from the report CONTEXT (derived artifacts only): title, summary, what we analysed, sensor
understanding, data trust, monitoring, suspicious rows (when the run has them), the top diagnoses (one slide each),
fault patterns (when present), human decisions and log integrity, data-flow record, assessor verdict with the learning
curve, evaluation (when labels exist) and open points / next steps.

Design: one accent colour, one type scale, fixed margins. Charts are native PowerPoint charts (python-pptx chart API,
with the embedded workbook) and tables are native tables, so the deck stays editable. Every text box is filled against
a line budget computed from its size and font: text that does not fit is cut at a sentence boundary with an ellipsis
and the slide says that the full wording is in the report. Speaker notes carry the evidence ids behind each slide.
"""
from __future__ import annotations

import io
import math
import re
from pathlib import Path
from typing import Any, Optional, Sequence

from ..config import Settings, get_settings
from ..workspace import Workspace
from .export_common import ELLIPSIS, cause_split, export_context, export_path, fit_text, headline_numbers, ids_in, interesting_signals, log_export, next_steps, open_points, role_mix, top_diagnoses, write_atomic_bytes
from .i18n import normalize_lang
from .report import _num, _pct, _thousands

SLIDE_W, SLIDE_H = 13.333, 7.5  # inches (16:9)
MARGIN = 0.6
CONTENT_TOP = 1.62
CONTENT_BOTTOM = 6.62
FONT_NAME = "Calibri"
CHAR_EM = 0.45  # average glyph width as a share of the font size (Calibri, mixed text; measured about 0.40), kept conservative
CELL_CHAR_EM = 0.42  # table cells hold short tokens (ids, numbers, labels): measured about 0.38
LINE_EM = 1.2
INSET = 0.1  # text-frame inset, inches, each side
ACCENT, INK, MUTED, GRID, SOFT, WHITE = "1A4D8F", "1F2937", "6B7280", "D9DDE3", "EEF4FB", "FFFFFF"
PASS, WARN, FAIL, WARN_BG = "2E7D32", "ED6C02", "C62828", "FFF4E5"
MAX_SUSPICIOUS_PER_TABLE = 13


# ----------------------------------------------------------------------------- text budget
def chars_per_line(width_in: float, size_pt: float, indent_in: float = 0.0) -> int:
    return max(6, int((width_in - 2 * INSET - indent_in) * 72.0 / (size_pt * CHAR_EM)))


def lines_available(height_in: float, size_pt: float) -> int:
    return max(1, int((height_in - 2 * INSET * 0.6) * 72.0 / (size_pt * LINE_EM)))


def text_budget(width_in: float, height_in: float, size_pt: float) -> int:
    """Characters a box of this size holds at this font size (what the tests check every text frame against)."""
    return chars_per_line(width_in, size_pt) * lines_available(height_in, size_pt)


def fit_paragraphs(paras: Sequence[dict[str, Any]], width_in: float, height_in: float) -> tuple[list[dict[str, Any]], bool]:
    """Keep the paragraphs that fit into the box, line by line; the last one that does not fit is cut at a sentence
    boundary (then at a word) with an ellipsis. Returns (paragraphs, was anything shortened or dropped)."""
    room = (height_in - 2 * INSET * 0.6) * 72.0
    out: list[dict[str, Any]] = []
    cut = False
    for p in paras:
        text = re.sub(r"\s+", " ", str(p.get("text") or "")).strip()
        if not text:
            continue
        size = float(p.get("size") or 12)
        cpl = chars_per_line(width_in, size, 0.24 if p.get("bullet") else 0.0)
        line_h = size * LINE_EM
        before = float(p.get("before") or 0.0) if out else 0.0
        after = float(p.get("after") if p.get("after") is not None else size * 0.35)
        lines_free = int((room - before) // line_h)
        if lines_free < 1:
            cut = True
            break
        need = _lines(text, cpl)
        if need > lines_free:
            cut = True
            if p.get("optional") and out:
                continue
            text, _ = fit_text(text, int(lines_free * cpl / WRAP_LOSS) - 1)
            if text:
                out.append({**p, "text": text})
            break  # the box is full: whatever follows is dropped
        out.append({**p, "text": text})
        room -= before + need * line_h + after
    return out, cut


WRAP_LOSS = 1.08  # words wrap before the line is full: count 8 % more characters than there are


def _lines(text: str, cpl: int) -> int:
    return max(1, math.ceil(len(text) * WRAP_LOSS / cpl))


def cell_capacity(col_w_in: float, row_h_in: float, size_pt: float) -> int:
    lines = max(1, int((row_h_in - 0.08) * 72.0 / (size_pt * LINE_EM)))
    return max(4, int((col_w_in - 0.16) * 72.0 / (size_pt * CELL_CHAR_EM))) * lines


# ----------------------------------------------------------------------------- deck
class _Deck:
    def __init__(self, ctx: dict[str, Any]):
        from pptx import Presentation
        from pptx.util import Inches

        self.ctx = ctx
        self.t = ctx["t"]
        self.lang = ctx["lang"]
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.prs.core_properties.title = f"{self.t('title')} — {ctx['run_id']}"
        self.prs.core_properties.author = "Trustworthy Process Monitor"
        self.prs.core_properties.subject = self.t("subtitle")
        self.prs.core_properties.language = self.lang
        self._cut: dict[int, bool] = {}
        self.stats = {"charts": 0, "tables": 0, "shortened": 0}

    # ------------------------------------------------------------------ primitives
    @staticmethod
    def rgb(hex6: str):
        from pptx.dml.color import RGBColor

        return RGBColor.from_string(hex6)

    def rect(self, slide, x: float, y: float, w: float, h: float, fill: Optional[str] = SOFT, line: Optional[str] = None, rounded: bool = False, name: str = "tpm-shape"):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt

        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.name = name
        shp.shadow.inherit = False
        if rounded:
            try:
                shp.adjustments[0] = 0.08
            except Exception:
                pass
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = self.rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        return shp

    def text(self, slide, x: float, y: float, w: float, h: float, paras: Sequence[dict[str, Any]], anchor: str = "top", align: str = "left", shape=None, role: str = "text") -> bool:
        """Add a text box (or fill `shape`) with paragraphs {"text", "size", "bold", "italic", "color", "bullet":
        None | "•" | "1", "after", "before", "optional"}; everything is fitted to the box first. Returns True when
        text was shortened."""
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Emu, Inches, Pt

        fitted, cut = fit_paragraphs(paras, w, h)
        box = shape if shape is not None else slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        min_size = min([float(p.get("size") or 12) for p in fitted] or [12.0])
        box.name = f"tpm-{role};budget={text_budget(w, h, min_size)}"
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(INSET)
        tf.margin_top = tf.margin_bottom = Inches(INSET * 0.6)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        number = 0
        for i, p in enumerate(fitted):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[p.get("align") or align]
            size = float(p.get("size") or 12)
            para.line_spacing = 1.0
            para.space_after = Pt(float(p.get("after") if p.get("after") is not None else size * 0.35))
            if p.get("before") and i:
                para.space_before = Pt(float(p["before"]))
            if p.get("bullet"):
                pPr = para._p.get_or_add_pPr()
                pPr.set("marL", str(int(Emu(Inches(0.24)))))
                pPr.set("indent", str(-int(Emu(Inches(0.24)))))
                if p["bullet"] == "1":
                    number += 1
                    bu = OxmlElement("a:buAutoNum")
                    bu.set("type", "arabicPeriod")
                    if number > 1:
                        bu.set("startAt", str(number))
                else:
                    bu = OxmlElement("a:buChar")
                    bu.set("char", "•")
                pPr.append(bu)
            run = para.add_run()
            run.text = p["text"]
            f = run.font
            f.name = FONT_NAME
            f.size = Pt(size)
            f.bold = bool(p.get("bold"))
            f.italic = bool(p.get("italic"))
            f.color.rgb = self.rgb(p.get("color") or INK)
        if cut:
            self._cut[id(slide)] = True
        return cut

    def chip(self, slide, x: float, y: float, label: str, color: str = ACCENT, w: Optional[float] = None) -> float:
        """Small filled label; returns its width so that chips can be laid out in a row."""
        size = 11.0
        label, _ = fit_text(label, 44)
        w = w or min(4.2, 0.36 + len(label) * size * CHAR_EM / 72.0 * 1.08)
        shp = self.rect(slide, x, y, w, 0.34, fill=color, rounded=True)
        self.text(slide, x, y, w, 0.34, [{"text": label, "size": size, "bold": True, "color": WHITE, "after": 0}], anchor="middle", align="center", shape=shp, role="chip")
        return w

    def table(self, slide, x: float, y: float, w: float, header: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[float], size: float = 10.5, row_h: float = 0.36, max_lines: int = 1, header_size: Optional[float] = None, header_h: float = 0.36):
        """Native table. Every cell is cut to `max_lines` lines of its column, so the rows keep their height and the
        table never grows past the slide."""
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt

        total = float(sum(widths))
        col_w = [w * c / total for c in widths]
        h_size = header_size or size
        body_h = max(row_h, (max_lines * size * LINE_EM) / 72.0 + 0.085)
        n = len(rows) + 1
        gf = slide.shapes.add_table(n, len(header), Inches(x), Inches(y), Inches(w), Inches(header_h + body_h * len(rows)))
        gf.name = "tpm-table"
        tbl = gf.table
        tbl.first_row = True
        tbl.horz_banding = False
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Inches(cw)
        tbl.rows[0].height = Inches(header_h)
        for i in range(1, n):
            tbl.rows[i].height = Inches(body_h)
        cut_any = False
        for i in range(n):
            for j in range(len(header)):
                cell = tbl.cell(i, j)
                raw = header[j] if i == 0 else rows[i - 1][j]
                fs = h_size if i == 0 else size
                cap = cell_capacity(col_w[j], header_h if i == 0 else body_h, fs)
                txt, cut = fit_text("" if raw is None else raw, cap)
                cut_any = cut_any or (cut and i > 0)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb(ACCENT if i == 0 else (WHITE if i % 2 else "F6F8FB"))
                tf = cell.text_frame
                tf.word_wrap = True
                para = tf.paragraphs[0]
                run = para.add_run()
                run.text = txt
                run.font.name = FONT_NAME
                run.font.size = Pt(fs)
                run.font.bold = i == 0
                run.font.color.rgb = self.rgb(WHITE if i == 0 else INK)
        if cut_any:
            self._cut[id(slide)] = True
        self.stats["tables"] += 1
        return gf

    # ------------------------------------------------------------------ charts (native, editable)
    def _style_chart(self, chart, legend: bool = False, size: float = 10.0) -> None:
        from pptx.enum.chart import XL_LEGEND_POSITION
        from pptx.util import Pt

        chart.has_title = False
        chart.font.size = Pt(size)
        chart.font.name = FONT_NAME
        chart.font.color.rgb = self.rgb(INK)
        chart.has_legend = legend
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(size)

    def _style_axes(self, chart, value_max: Optional[float] = None, value_fmt: Optional[str] = None, gridlines: bool = True) -> None:
        from pptx.enum.chart import XL_TICK_MARK
        from pptx.util import Pt

        va, ca = chart.value_axis, chart.category_axis
        va.has_major_gridlines = gridlines
        if gridlines:
            va.major_gridlines.format.line.color.rgb = self.rgb(GRID)
            va.major_gridlines.format.line.width = Pt(0.5)
        va.format.line.fill.background()
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.tick_labels.font.size = Pt(9)
        va.tick_labels.font.color.rgb = self.rgb(MUTED)
        if value_max is not None:
            va.maximum_scale = value_max
            va.minimum_scale = 0
        if value_fmt:
            va.tick_labels.number_format = value_fmt
            va.tick_labels.number_format_is_linked = False
        ca.format.line.color.rgb = self.rgb(GRID)
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.tick_labels.font.size = Pt(10)

    def bar_chart(self, slide, x: float, y: float, w: float, h: float, items: Sequence[tuple[str, float]], series_name: str, percent: bool = False, color: str = ACCENT, value_max: Optional[float] = None):
        """Horizontal bars, the first item on top."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
        from pptx.util import Inches, Pt

        items = [(fit_text(k, 34)[0], float(v)) for k, v in items][::-1]  # bar charts plot bottom-up
        cd = CategoryChartData()
        cd.categories = [k for k, _ in items]
        cd.add_series(series_name, [v for _, v in items])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "tpm-chart"
        chart = gf.chart
        self._style_chart(chart)
        self._style_axes(chart, value_max=value_max, value_fmt="0%" if percent else None)
        plot = chart.plots[0]
        plot.gap_width = 55
        plot.vary_by_categories = False
        ser = plot.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = self.rgb(color)
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(9)
        dl.font.color.rgb = self.rgb(MUTED)
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.number_format = "0%" if percent else "General"
        dl.number_format_is_linked = False
        self.stats["charts"] += 1
        return gf

    def stacked_chart(self, slide, x: float, y: float, w: float, h: float, categories: Sequence[str], series: Sequence[tuple[str, Sequence[float], str]]):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        cd = CategoryChartData()
        cd.categories = [fit_text(c, 28)[0] for c in categories][::-1]
        for name, vals, _ in series:
            cd.add_series(name, list(vals)[::-1])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "tpm-chart"
        chart = gf.chart
        self._style_chart(chart, legend=True)
        self._style_axes(chart)
        plot = chart.plots[0]
        plot.gap_width = 45
        plot.overlap = 100
        for ser, (_, _, color) in zip(plot.series, series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = self.rgb(color)
        self.stats["charts"] += 1
        return gf

    def column_chart(self, slide, x: float, y: float, w: float, h: float, labels: Sequence[str], values: Sequence[float], series_name: str, red_below: Optional[float] = None, flags: Optional[Sequence[bool]] = None):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches, Pt

        cd = CategoryChartData()
        cd.categories = list(labels)
        cd.add_series(series_name, list(values))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "tpm-chart"
        chart = gf.chart
        self._style_chart(chart)
        self._style_axes(chart, value_max=1.0)
        chart.category_axis.tick_labels.font.size = Pt(8)
        self._tick_skip(chart, max(1, len(labels) // 8))
        plot = chart.plots[0]
        plot.gap_width = 30
        plot.vary_by_categories = False
        ser = plot.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = self.rgb(ACCENT)
        for i, v in enumerate(values):
            bad = (flags[i] if flags is not None else False) or (red_below is not None and v < red_below)
            if bad:
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = self.rgb(FAIL)
        self.stats["charts"] += 1
        return gf

    def line_chart(self, slide, x: float, y: float, w: float, h: float, labels: Sequence[str], series: Sequence[tuple[str, Sequence[Optional[float]], str, float, bool]]):
        """series: (name, values, colour, width pt, dashed)."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
        from pptx.enum.dml import MSO_LINE
        from pptx.util import Inches, Pt

        cd = CategoryChartData()
        cd.categories = list(labels)
        for name, vals, *_ in series:
            cd.add_series(name, list(vals))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "tpm-chart"
        chart = gf.chart
        self._style_chart(chart, legend=True)
        self._style_axes(chart)
        chart.category_axis.tick_labels.font.size = Pt(9)
        self._tick_skip(chart, max(1, len(labels) // 6))
        for ser, (_, _, color, width, dashed) in zip(chart.plots[0].series, series):
            ser.smooth = False
            ser.marker.style = XL_MARKER_STYLE.NONE
            ser.format.line.color.rgb = self.rgb(color)
            ser.format.line.width = Pt(width)
            if dashed:
                ser.format.line.dash_style = MSO_LINE.DASH
        self.stats["charts"] += 1
        return gf

    def xy_chart(self, slide, x: float, y: float, w: float, h: float, points: Sequence[tuple[float, float]], series_name: str):
        from pptx.chart.data import XyChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE, XL_TICK_MARK
        from pptx.util import Inches, Pt

        cd = XyChartData()
        s = cd.add_series(series_name)
        for a, b in points:
            s.add_data_point(float(a), float(b))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "tpm-chart"
        chart = gf.chart
        self._style_chart(chart)
        for ax in (chart.value_axis, chart.category_axis):
            ax.format.line.color.rgb = self.rgb(GRID)
            ax.major_tick_mark = XL_TICK_MARK.NONE
            ax.tick_labels.font.size = Pt(9)
            ax.tick_labels.font.color.rgb = self.rgb(MUTED)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = self.rgb(GRID)
        chart.category_axis.has_major_gridlines = False
        ys = [float(b) for _, b in points]
        xs = [float(a) for a, _ in points]
        chart.category_axis.minimum_scale = 0 if 0 <= min(xs) <= 1 else min(xs)
        chart.category_axis.maximum_scale = max(xs)
        pad = (max(ys) - min(ys)) * 0.25 or 0.05
        chart.value_axis.minimum_scale = round(max(0.0, min(ys) - pad), 2) if min(ys) >= 0 else round(min(ys) - pad, 2)
        chart.value_axis.maximum_scale = round(max(ys) + pad, 2)
        ser = chart.plots[0].series[0]
        ser.smooth = False
        ser.format.line.color.rgb = self.rgb(ACCENT)
        ser.format.line.width = Pt(2.25)
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 7
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = self.rgb(ACCENT)
        self.stats["charts"] += 1
        return gf

    @staticmethod
    def _tick_skip(chart, every: int) -> None:
        """Label every n-th category (long series): c:tickLblSkip / c:tickMarkSkip in schema order."""
        from pptx.oxml.ns import qn
        from pptx.oxml.xmlchemy import OxmlElement

        if every <= 1:
            return
        ax = chart.category_axis._element
        for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
            for old in ax.findall(qn(tag)):
                ax.remove(old)
        anchor = ax.find(qn("c:noMultiLvlLbl"))
        for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
            el = OxmlElement(tag)
            el.set("val", str(int(every)))
            if anchor is not None:
                anchor.addprevious(el)
            else:
                ax.append(el)

    # ------------------------------------------------------------------ slide scaffolding
    def new_slide(self, title: str, kicker: str = ""):
        slide = self.prs.slides.add_slide(self.blank)
        self.rect(slide, MARGIN, 0.5, 0.55, 0.07, fill=ACCENT)
        if kicker:
            self.text(slide, MARGIN - INSET, 0.6, SLIDE_W - 2 * MARGIN, 0.34, [{"text": kicker.upper(), "size": 11, "bold": True, "color": ACCENT, "after": 0}], role="kicker")
        self.text(slide, MARGIN - INSET, 0.9, SLIDE_W - 2 * MARGIN + 2 * INSET, 0.62, [{"text": title, "size": 26, "bold": True, "color": INK, "after": 0}], anchor="middle", role="title")
        return slide

    def finish(self, slide, number: int, evidence: Sequence[str], section: str) -> None:
        """Footer (run id, slide number field), the shortened-text note and the speaker notes."""
        from pptx.enum.text import PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches, Pt

        t = self.t
        if self._cut.get(id(slide)):
            self.stats["shortened"] += 1
            self.text(slide, MARGIN - INSET, 6.66, 9.0, 0.3, [{"text": t("px_shortened"), "size": 9, "italic": True, "color": MUTED, "after": 0}], role="note")
        self.rect(slide, MARGIN, 6.98, SLIDE_W - 2 * MARGIN, 0.01, fill=GRID)
        self.text(slide, MARGIN - INSET, 7.02, 9.5, 0.3, [{"text": f"{t('px_footer', run_id=self.ctx['run_id'])} · {self.ctx['generated_at'][:10]}", "size": 9, "color": MUTED, "after": 0}], role="footer")
        box = slide.shapes.add_textbox(Inches(SLIDE_W - MARGIN - 1.2 + INSET), Inches(7.02), Inches(1.2), Inches(0.3))
        box.name = "tpm-slide-number;budget=6"
        para = box.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT
        fld = OxmlElement("a:fld")
        fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
        fld.set("type", "slidenum")
        rpr = OxmlElement("a:rPr")
        rpr.set("lang", "en-US")
        rpr.set("sz", "900")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", MUTED)
        fill.append(clr)
        rpr.append(fill)
        fld.append(rpr)
        tx = OxmlElement("a:t")
        tx.text = str(number)
        fld.append(tx)
        para._p.append(fld)
        self.notes(slide, evidence, section)

    def notes(self, slide, evidence: Sequence[str], section: str) -> None:
        t = self.t
        ids = [i for i in dict.fromkeys(str(x) for x in evidence if x)][:40]
        lines = [t("px_notes_evidence", ids=", ".join(ids)) if ids else t("px_notes_none"), t("px_notes_source", section=section, run_id=self.ctx["run_id"])]
        if self._cut.get(id(slide)):
            lines.append(t("px_shortened"))
        slide.notes_slide.notes_text_frame.text = "\n".join(lines)

    def heading(self, slide, x: float, y: float, w: float, label: str) -> float:
        """Column heading; returns the y below it."""
        self.text(slide, x - INSET, y, w + 2 * INSET, 0.36, [{"text": label, "size": 13, "bold": True, "color": ACCENT, "after": 0}], anchor="middle", role="heading")
        return y + 0.42

    # ------------------------------------------------------------------ slides
    def slide_title(self) -> None:
        ctx, t = self.ctx, self.t
        df = ctx["dataflow"]
        slide = self.prs.slides.add_slide(self.blank)
        self.rect(slide, 0, 0, 4.4, SLIDE_H, fill=ACCENT)
        self.text(slide, 0.5, 0.55, 3.5, 0.9, [{"text": "TRUSTWORTHY PROCESS MONITOR", "size": 13, "bold": True, "color": WHITE, "after": 0}], role="brand")
        self.text(slide, 0.5, 5.9, 3.5, 1.1, [{"text": t("pdf_no_raw"), "size": 10, "color": "D6E2F3", "after": 0}], anchor="bottom", role="brand-note")
        x = 5.0
        w = SLIDE_W - x - MARGIN
        self.text(slide, x, 1.1, w, 1.5, [{"text": t("title"), "size": 34, "bold": True, "color": INK, "after": 0}], anchor="bottom", role="title")
        self.text(slide, x, 2.65, w, 0.8, [{"text": t("subtitle"), "size": 16, "color": MUTED, "after": 0}], role="subtitle")
        self.rect(slide, x + INSET, 3.55, 0.9, 0.07, fill=ACCENT)
        meta = [(t("run_id"), ctx["run_id"]), (t("source"), ctx["source_name"]), (t("generated_at"), f"{ctx['generated_at']} UTC"), (t("profile"), df.get("profile")), (t("s8_local_model"), f"{df.get('local_provider')} / {df.get('local_model')}"), (t("language"), ctx["lang_name"])]
        self.text(slide, x, 3.8, w, 2.6, [{"text": f"{k}: {v}", "size": 14, "color": INK, "after": 5} for k, v in meta if v], role="meta")
        self.notes(slide, [], t("section_overview"))

    def slide_summary(self, n: int) -> None:
        ctx, t = self.ctx, self.t
        slide = self.new_slide(t("px_summary"), t("section_overview"))
        left_w = 6.6
        llm = ctx.get("llm")
        body_h = (CONTENT_BOTTOM - CONTENT_TOP) if not llm else 2.9
        self.text(slide, MARGIN - INSET, CONTENT_TOP, left_w, body_h, [{"text": s, "size": 13, "bullet": "•", "after": 6} for s in ctx["overview"][:6]], role="body")
        if llm:
            y = CONTENT_TOP + body_h + 0.1
            h = CONTENT_BOTTOM - y
            self.rect(slide, MARGIN, y, left_w - INSET, h, fill="FFFBEB", line=WARN)
            paras = [{"text": f"{t('section_llm')} — {llm.get('source') or 'llm'}", "size": 10.5, "bold": True, "color": WARN, "after": 3}]
            paras += [{"text": p, "size": 11.5, "after": 3} for p in (llm.get("summary") or [])[:2]]
            self.text(slide, MARGIN, y, left_w - INSET, h, paras, role="llm")
        tiles = headline_numbers(ctx)[:8]
        tx, ty, tw, th, gap = MARGIN + left_w + 0.35, CONTENT_TOP + 0.05, 2.55, 1.1, 0.14
        for k, (label, value) in enumerate(tiles):
            cx, cy = tx + (k % 2) * (tw + gap), ty + (k // 2) * (th + gap)
            self.rect(slide, cx, cy, tw, th, fill=SOFT)
            self.text(slide, cx, cy + 0.08, tw, 0.6, [{"text": value, "size": 22, "bold": True, "color": ACCENT, "after": 0}], anchor="middle", align="center", role="tile-number")
            self.text(slide, cx, cy + 0.66, tw, 0.36, [{"text": label, "size": 10.5, "color": MUTED, "after": 0}], anchor="middle", align="center", role="tile-label")
        self.finish(slide, n, ids_in(ctx["overview"], (llm or {}).get("summary")), t("section_overview"))

    def slide_analysed(self, n: int) -> None:
        ctx, t, lang = self.ctx, self.t, self.lang
        slide = self.new_slide(t("px_analysed"), t("section_1"))
        ds = ctx.get("dataset") or {}
        rows = [[t("source"), f"{ctx['source_name']} ({ds.get('format') or t('na_short')})"], [t("rows"), _thousands(ds.get("n_rows"), lang) if ds else t("na_short")], [t("columns"), ds.get("n_cols", t("na_short"))], [t("signals"), f"{ctx.get('n_signals_total', 0)}"], [t("groups"), f"{_thousands(ds.get('n_groups'), lang)} ({ds.get('grouping_method')})" if ds else t("na_short")], [t("batches"), _thousands(ctx["quality"].get("n_batches", 0), lang)], [t("s1_time_column"), ds.get("time_column") or t("none")], [t("s1_sample_period"), ds.get("sample_period_text") or t("na_short")], [t("s1_label_columns").split(" (")[0], ", ".join(ds.get("label_columns") or []) or t("none")]]
        self.table(slide, MARGIN, CONTENT_TOP + 0.05, 5.9, [t("key"), t("value")], rows, [42, 58], size=11.5, row_h=0.4)
        x = MARGIN + 6.35
        w = SLIDE_W - MARGIN - x
        y = self.heading(slide, x, CONTENT_TOP - 0.05, w, t("px_kind_of_data"))
        dom = [(k, float(str(v).replace("%", "").strip() or 0) / 100.0) for k, v in (ds.get("domain_items") or []) if re.match(r"^\s*\d+(\.\d+)?\s*%\s*$", str(v))]
        if dom:
            self.bar_chart(slide, x - 0.1, y, w + 0.1, 1.9, [(k.replace("_", " "), v) for k, v in dom[:5]], t("s1_domain"), percent=True, value_max=1.0)
            y += 2.0
        y = self.heading(slide, x, y, w, t("assumptions"))
        assumptions = list(ds.get("assumptions") or []) + list(ctx["detect"].get("baseline_assumptions") or [])
        ev = list(ds.get("evidence_ids") or []) + [i for inf in (ds.get("inferences") or [])[:4] for i in (inf.get("evidence_ids") or [])[:2]]
        self.text(slide, x - INSET, y, w + INSET, CONTENT_BOTTOM - y, [{"text": a, "size": 11.5, "bullet": "•", "after": 4} for a in assumptions[:4]] or [{"text": t("not_available"), "size": 11.5, "color": MUTED}], role="body")
        self.finish(slide, n, ev + [inf.get("id") for inf in (ds.get("inferences") or [])[:4]], t("section_1"))

    def slide_sensors(self, n: int) -> None:
        ctx, t = self.ctx, self.t
        slide = self.new_slide(t("px_sensors"), t("section_1"))
        lw = 5.4
        y = self.heading(slide, MARGIN, CONTENT_TOP - 0.05, lw, t("px_role_mix"))
        mix = role_mix(ctx)
        if mix:
            self.bar_chart(slide, MARGIN - 0.1, y, lw + 0.1, CONTENT_BOTTOM - y, [(k, v) for k, v in mix], t("signals"))
        else:
            self.text(slide, MARGIN - INSET, y, lw, 0.6, [{"text": t("s1_no_signals"), "size": 12, "color": MUTED}], role="body")
        x = MARGIN + lw + 0.5
        w = SLIDE_W - MARGIN - x
        y = self.heading(slide, x, CONTENT_TOP - 0.05, w, t("px_interesting_signals"))
        picks = interesting_signals(ctx, 4)
        ev: list[str] = []
        if picks:
            h_each = (CONTENT_BOTTOM - y) / len(picks)
            for k, s in enumerate(picks):
                head = f"{s['id']} — {s.get('role_label')} ({_pct(s.get('confidence'))})" + (f" · {t('px_leads_flags', n=s['n_led'])}" if s.get("n_led") else (f" · {s['instrument']}?" if s.get("instrument") else ""))
                self.text(slide, x - INSET, y + k * h_each, w + INSET, h_each - 0.05, [{"text": head, "size": 12.5, "bold": True, "after": 2}, {"text": s.get("sentence") or t("not_available"), "size": 11, "color": "374151", "after": 0}], role="body")
                ev += s.get("evidence_ids") or []
        else:
            self.text(slide, x - INSET, y, w, 0.6, [{"text": t("s1_no_signals"), "size": 12, "color": MUTED}], role="body")
        self.finish(slide, n, ev, t("section_1"))

    def slide_trust(self, n: int) -> None:
        ctx, t, lang = self.ctx, self.t, self.lang
        q = ctx["quality"]
        slide = self.new_slide(t("px_trust"), t("section_2"))
        lw = 6.3
        series = q.get("trust_series") or []
        y = CONTENT_TOP - 0.05
        if series:
            binned = (q.get("n_trust_total") or 0) > len(series)
            y = self.heading(slide, MARGIN, y, lw, t("px_trust_by_block") if binned else t("px_trust_by_batch"))
            self.column_chart(slide, MARGIN - 0.1, y, lw + 0.1, 2.15, [str(p.get("label")) for p in series], [float(p.get("score") or 0) for p in series], t("trust_score"), flags=[not p.get("trusted", True) for p in series])
            y += 2.2
        cats = q.get("by_category") or []
        if cats:
            y = self.heading(slide, MARGIN, y, lw, t("px_checks_by_category"))
            self.stacked_chart(slide, MARGIN - 0.1, y, lw + 0.1, CONTENT_BOTTOM - y, [c for c, _ in cats], [(t("pass"), [int(v.get("pass", 0)) for _, v in cats], PASS), (t("warn"), [int(v.get("warn", 0)) for _, v in cats], WARN), (t("fail"), [int(v.get("fail", 0)) for _, v in cats], FAIL)])
        elif not series:
            self.text(slide, MARGIN - INSET, y, lw, 0.6, [{"text": t("s2_no_checks"), "size": 12, "color": MUTED}], role="body")
        x = MARGIN + lw + 0.5
        w = SLIDE_W - MARGIN - x
        verdict = t("overview_quality", n_checks=_thousands(q.get("n_checks", 0), lang), n_pass=_thousands(q.get("n_pass", 0), lang), n_warn=_thousands(q.get("n_warn", 0), lang), n_fail=_thousands(q.get("n_fail", 0), lang), n_untrusted=_thousands(q.get("n_untrusted_total", 0), lang), n_batches=_thousands(q.get("n_batches", 0), lang))
        self.text(slide, x - INSET, CONTENT_TOP - 0.05, w + INSET, 0.95, [{"text": verdict, "size": 12.5, "after": 0}], role="body")
        y2 = self.heading(slide, x, CONTENT_TOP + 0.95, w, t("px_top_findings"))
        failed = (q.get("failed") or [])[:3]
        self.text(slide, x - INSET, y2, w + INSET, 1.95, [{"text": c.get("statement"), "size": 11, "bullet": "•", "after": 4} for c in failed] or [{"text": t("s2_all_pass") if q.get("n_checks") else t("s2_no_checks"), "size": 11.5, "color": MUTED}], role="body")
        y3 = self.heading(slide, x, y2 + 2.0, w, t("px_sensor_vs_process"))
        split = cause_split(ctx)
        self.text(slide, x - INSET, y3, w + INSET, CONTENT_BOTTOM - y3, [{"text": t("px_sensor_vs_process_text", n_sensor=_thousands(split["sensor"], lang), n_process=_thousands(split["process"], lang), n_other=_thousands(split["other"], lang)), "size": 11, "after": 0}], role="body")
        self.finish(slide, n, [c.get("check_id") for c in failed] + [e for c in failed for e in (c.get("evidence_ids") or [])[:2]], t("section_2"))

    def slide_monitoring(self, n: int) -> None:
        ctx, t, lang = self.ctx, self.t, self.lang
        det = ctx["detect"]
        slide = self.new_slide(t("px_monitoring"), t("section_3"))
        tl = det.get("timelines") or {}
        lines = [x for x in (tl.get("timelines") or []) if x.get("values")]
        lw = 7.6
        ev: list[str] = []
        if lines:
            best = max(lines, key=lambda x: (int(x.get("n_flags") or 0) > 0, float(x.get("max") or 0)))
            vals = [float(v) for v in best["values"]]
            th = tl.get("threshold")
            nv = len(vals)
            r0, r1 = (best.get("x_labels") or ("0", str(nv)))[0], (best.get("x_labels") or ("0", str(nv)))[-1]
            try:
                a, b = float(r0), float(r1)
                labels = [str(int(a + (b - a) * i / max(1, nv - 1))) for i in range(nv)]
            except ValueError:
                labels = [str(i) for i in range(nv)]
            flagged: list[Optional[float]] = [None] * nv
            for s0, s1 in best.get("spans") or []:
                for i in range(max(0, int(s0)), min(nv - 1, int(s1)) + 1):
                    flagged[i] = vals[i]
            series = [(t("px_score_series"), vals, ACCENT, 1.75, False)]
            if any(v is not None for v in flagged):
                series.append((t("s3_flags"), flagged, FAIL, 3.0, False))
            if th is not None:
                series.append((t("threshold"), [float(th)] * nv, WARN, 1.25, True))
            y = self.heading(slide, MARGIN, CONTENT_TOP - 0.05, lw, t("px_score_timeline", group=best.get("group")))
            self.line_chart(slide, MARGIN - 0.1, y, lw + 0.1, CONTENT_BOTTOM - y, labels, series)
        else:
            self.text(slide, MARGIN - INSET, CONTENT_TOP, lw, 0.6, [{"text": t("s3_no_scores"), "size": 12, "color": MUTED}], role="body")
        x = MARGIN + lw + 0.45
        w = SLIDE_W - MARGIN - x
        gs = det.get("group_summary") or {}
        lead = t("px_groups_over", n_over=_thousands(gs.get("n_over", 0), lang), n_groups=_thousands(gs.get("n_groups", 0), lang)) if gs.get("n_groups") else ""
        n_groups_flagged = len({f.get("group_id") for f in det.get("flags") or []})
        paras = [{"text": lead, "size": 15, "bold": True, "color": ACCENT, "after": 6}] if lead else []
        paras.append({"text": t("overview_detect", n_flags=_thousands(det.get("n_flags", 0), lang), n_groups_flagged=_thousands(gs.get("n_over", n_groups_flagged), lang), n_patterns=len(det.get("patterns") or [])), "size": 12, "after": 0})
        self.text(slide, x - INSET, CONTENT_TOP - 0.05, w + INSET, 1.5, paras, role="body")
        kinds = (det.get("flags_by_kind") or [])[:4]
        if kinds:
            self.table(slide, x, CONTENT_TOP + 1.55, w, [t("kind"), t("count")], [[k, _thousands(v, lang)] for k, v in kinds], [68, 32], size=11, row_h=0.34)
        base = dict(det.get("baseline_items") or [])
        y4 = CONTENT_TOP + 1.55 + 0.36 + 0.34 * len(kinds) + 0.25
        if y4 < CONTENT_BOTTOM - 0.6:
            note = [a for a in (det.get("baseline_assumptions") or [])[:1]]
            txt = (f"{t('s3_baseline')}: {base.get('strategy') or base.get('method') or t('na_short')}" + (f" ({t('confidence').lower()} {_pct(base.get('confidence') if base.get('confidence') is not None else base.get('score'))})" if (base.get("confidence") is not None or base.get("score") is not None) else "") + ".") if base else ""
            self.text(slide, x - INSET, y4, w + INSET, CONTENT_BOTTOM - y4, [{"text": txt, "size": 11, "bold": True, "after": 3}] + [{"text": a, "size": 10.5, "color": "374151", "after": 0, "optional": True} for a in note], role="body")
            ev += ids_in(base.get("evidence_ids"))
        top_flags = (det.get("flags") or [])[:6]
        self.finish(slide, n, ev + [f.get("id") for f in top_flags] + [e for f in top_flags[:3] for e in (f.get("evidence_ids") or [])[:2]], t("section_3"))

    def slide_suspicious(self, n: int) -> None:
        ctx, t, lang = self.ctx, self.t, self.lang
        su = ctx["suspicious"]
        title = t("section_suspicious") + (f" ({_thousands(su['n_rows'], lang)})" if su.get("n_rows") else "")
        slide = self.new_slide(title, t("section_3"))
        # the wording is the message of this slide: it is never cut, the banner grows and the type steps down instead
        full_w = SLIDE_W - 2 * MARGIN
        size, lines = next(((sz, n) for sz in (17.0, 15.0, 13.5, 12.0) for n in (_lines(su["wording"], chars_per_line(full_w - 0.2, sz)),) if n <= (1 if sz == 17.0 else 2 if sz >= 13.5 else 3)), (12.0, 3))
        banner_h = max(0.62, lines * size * LINE_EM / 72.0 + 0.3)
        self.rect(slide, MARGIN, CONTENT_TOP, full_w, banner_h, fill=WARN_BG)
        self.rect(slide, MARGIN, CONTENT_TOP, 0.08, banner_h, fill=WARN)
        self.text(slide, MARGIN + 0.15, CONTENT_TOP, full_w - 0.2, banner_h, [{"text": su["wording"], "size": size, "bold": True, "color": INK, "after": 0}], anchor="middle", role="wording")
        y = CONTENT_TOP + banner_h + 0.08
        sub = su.get("headline") or su.get("regime_text") or t("susp_intro")
        sub_h = 0.34 if _lines(sub, chars_per_line(full_w, 11.5)) <= 1 else 0.55
        self.text(slide, MARGIN - INSET, y, full_w, sub_h, [{"text": sub, "size": 11.5, "color": "374151", "after": 0}], role="body")
        y += sub_h + 0.05
        ev: list[str] = []
        if su["rows"]:
            brief = lambda r: " · ".join(" ".join(str(p) for p in (x.get("signal"), _num(x.get("deviation"), 1), x.get("direction_label") or x.get("direction")) if p not in (None, "")) for x in (r.get("signals") or [])[:2])  # noqa: E731
            rows = [[r.get("row_text"), r.get("group_id") or r.get("batch_id") or "", brief(r) + (" …" if len(r.get("signals") or []) > 2 else ""), r.get("sources_text")] for r in su["rows"]]
            for r in su["rows"]:
                ev += list(r.get("flag_ids") or [])[:1] + list(r.get("check_ids") or [])[:1] + list(r.get("evidence_ids") or [])[:1]
            header = [t("susp_row"), t("group"), t("susp_signals"), t("susp_found_by")]
            widths = [13, 9, 37, 41]
        else:
            rows = [[f.get("row_text"), f.get("group_id") or f.get("batch_id") or "", "; ".join(f.get("signals") or []), f.get("id")] for f in su["point_flags"]]
            ev += [f.get("id") for f in su["point_flags"]]
            header = [t("susp_row"), t("group"), t("responsible_signals"), t("id")]
            widths = [17, 13, 45, 25]
        rows = rows[: 2 * MAX_SUSPICIOUS_PER_TABLE]
        head_h = 0.32
        row_h = min(0.3, (CONTENT_BOTTOM - y - head_h) / max(1, min(len(rows), MAX_SUSPICIOUS_PER_TABLE)))
        size = 9.0 if row_h >= 0.235 else 8.5
        half = (SLIDE_W - 2 * MARGIN - 0.3) / 2
        if len(rows) > MAX_SUSPICIOUS_PER_TABLE:
            self.table(slide, MARGIN, y, half, header, rows[:MAX_SUSPICIOUS_PER_TABLE], widths, size=size, row_h=row_h, header_size=9.5, header_h=head_h)
            self.table(slide, MARGIN + half + 0.3, y, half, header, rows[MAX_SUSPICIOUS_PER_TABLE:], widths, size=size, row_h=row_h, header_size=9.5, header_h=head_h)
        elif rows:
            self.table(slide, MARGIN, y, SLIDE_W - 2 * MARGIN, header, rows, widths, size=10, row_h=row_h, header_h=head_h)
        else:
            self.text(slide, MARGIN - INSET, y, 8, 0.5, [{"text": t("susp_none_listed"), "size": 12, "color": MUTED}], role="body")
        self.finish(slide, n, ev, t("section_suspicious"))

    def slide_diagnosis(self, n: int, d: dict[str, Any], i: int, total: int) -> None:
        t = self.t
        slide = self.new_slide(fit_text(str(d.get("fault_type") or t("unknown")), 62)[0], f"{t('px_diagnosis', i=i, n=total)} · {d.get('id')} · {d.get('cause_label') or ''}".rstrip(" ·"))
        lw = 6.7
        y = self.heading(slide, MARGIN, CONTENT_TOP - 0.05, lw, t("px_what_happened"))
        self.text(slide, MARGIN - INSET, y, lw + INSET, 1.75, [{"text": d.get("summary") or t("not_available"), "size": 12, "after": 0}], role="body")
        y2 = self.heading(slide, MARGIN, y + 1.8, lw, t("steps"))
        steps = [s for s in (d.get("steps") or []) if not re.match(r"^\s*Model explanation\b", s)][:4] or (d.get("steps") or [])[:4]
        self.text(slide, MARGIN - INSET, y2, lw + INSET, CONTENT_BOTTOM - y2, [{"text": s, "size": 11, "bullet": "1", "after": 4} for s in steps] or [{"text": t("not_available"), "size": 11, "color": MUTED}], role="body")
        x = MARGIN + lw + 0.5
        w = SLIDE_W - MARGIN - x
        cx = x
        cx += self.chip(slide, cx, CONTENT_TOP, f"{t('confidence')} {_pct(d.get('confidence'))}") + 0.12
        crit = d.get("critique") or {}
        if crit:
            cx += self.chip(slide, cx, CONTENT_TOP, f"{t('critique')}: {d.get('verdict_label')}", PASS if crit.get("verdict") == "supported" else (FAIL if crit.get("verdict") == "rejected" else WARN)) + 0.12
        if d.get("human_status") and cx + 1.8 < SLIDE_W - MARGIN:
            self.chip(slide, cx, CONTENT_TOP, str(d["human_status"]), MUTED)
        ranked = (d.get("ranked") or d.get("ranked_signals") or [])[:6]
        y3 = self.heading(slide, x, CONTENT_TOP + 0.45, w, t("ranked_signals"))
        if ranked:
            ch = 0.55 + 0.3 * len(ranked)
            self.bar_chart(slide, x - 0.1, y3, w + 0.1, ch, [(f"{s.get('signal')}" + (f" ({s.get('direction')})" if s.get("direction") else ""), float(s.get("contribution") or 0)) for s in ranked], t("contribution"), percent=True)
            y3 += ch + 0.05
        y4 = self.heading(slide, x, y3, w, t("uncertain"))
        unc = list(d.get("uncertainty") or []) + list((crit.get("objections") or [])[:1])
        self.text(slide, x - INSET, y4, w + INSET, CONTENT_BOTTOM - y4, [{"text": u, "size": 11, "bullet": "•", "after": 4} for u in unc[:3]] or [{"text": t("unc_none"), "size": 11, "color": MUTED}], role="body")
        ev = [d.get("id")] + list(d.get("flag_ids") or [])[:4] + list(d.get("evidence_ids") or [])[:8] + ([d.get("pattern_id")] if d.get("pattern_id") else [])
        self.finish(slide, n, ev, t("section_4"))

    def slide_patterns(self, n: int) -> None:
        ctx, t = self.ctx, self.t
        pats = ctx["detect"].get("patterns") or []
        slide = self.new_slide(t("px_patterns"), t("section_3"))
        rows = [[p.get("id"), p.get("name_label"), p.get("n_events"), fit_text(", ".join(str(g) for g in (p.get("groups_affected") or [])), 40)[0], _pct(p.get("confidence")), p.get("description")] for p in pats[:6]]
        self.table(slide, MARGIN, CONTENT_TOP + 0.05, SLIDE_W - 2 * MARGIN, [t("pattern"), t("pattern_name"), t("n_events"), t("groups_affected"), t("confidence"), t("message")], rows, [11, 13, 8, 15, 10, 43], size=10.5, row_h=0.56, max_lines=2)
        y = CONTENT_TOP + 0.05 + 0.36 + 0.56 * len(rows) + 0.25
        ev = [p.get("id") for p in pats[:6]] + [e for p in pats[:6] for e in (p.get("evidence_ids") or [])[:2]]
        chain_d = next((d for d in ctx.get("diagnoses") or [] if d.get("propagation")), None)
        if chain_d and y < CONTENT_BOTTOM - 0.9:
            y = self.heading(slide, MARGIN, y, SLIDE_W - 2 * MARGIN, f"{t('propagation')} — {chain_d.get('id')}")
            steps = chain_d["propagation"][:4]
            self.text(slide, MARGIN - INSET, y, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - y, [{"text": f"{p.get('from_signal')} → {p.get('to_signal')}" + (f" ({t('lag')} {p.get('lag')})" if p.get("lag") is not None else "") + (f": {p.get('explanation')}" if p.get("explanation") else ""), "size": 11.5, "bullet": "•", "after": 4} for p in steps], role="body")
            ev += [chain_d.get("id")] + [e for p in steps for e in (p.get("evidence_ids") or [])[:1]]
        self.finish(slide, n, ev, t("section_3"))

    def slide_human(self, n: int) -> None:
        ctx, t, lang = self.ctx, self.t, self.lang
        log = ctx["log"]
        ok = bool((log.get("chain") or {}).get("ok"))
        slide = self.new_slide(t("px_human"), f"{t('section_5')} · {t('section_6')}")
        self.rect(slide, MARGIN, CONTENT_TOP, SLIDE_W - 2 * MARGIN, 0.6, fill="EAF5EA" if ok else WARN_BG)
        self.rect(slide, MARGIN, CONTENT_TOP, 0.08, 0.6, fill=PASS if ok else FAIL)
        self.text(slide, MARGIN + 0.15, CONTENT_TOP, SLIDE_W - 2 * MARGIN - 0.2, 0.6, [{"text": log.get("chain_text") or "", "size": 14, "bold": True, "color": PASS if ok else FAIL, "after": 0}], anchor="middle", role="banner")
        y = CONTENT_TOP + 0.8
        lw = 7.7
        human = ctx.get("human") or []
        y1 = self.heading(slide, MARGIN, y, lw, f"{t('section_5').split('. ', 1)[-1]} ({len(human)})")
        if human:
            rows = [[h.get("actor"), h.get("role"), h.get("action_label"), f"{h.get('object_type')} {h.get('object_id')}", h.get("note") or h.get("after")] for h in human[-6:]]
            self.table(slide, MARGIN, y1, lw, [t("actor"), t("role"), t("action"), t("object"), t("note")], rows, [14, 13, 16, 24, 33], size=10.5, row_h=0.5, max_lines=2)
        else:
            self.text(slide, MARGIN - INSET, y1, lw, 1.2, [{"text": t("s5_no_human"), "size": 12, "color": MUTED}], role="body")
        x = MARGIN + lw + 0.5
        w = SLIDE_W - MARGIN - x
        y2 = self.heading(slide, x, y, w, f"{t('s6_by_actor')} — {t('px_chain_entries', n=_thousands(log.get('total', 0), lang))}")
        actors = [(str(a).replace("system:", "").replace("llm:local:", "llm ").replace("llm:external:", "llm ext "), int(c)) for a, c in (log.get("by_actor") or [])[:7]]
        if actors:
            self.bar_chart(slide, x - 0.1, y2, w + 0.1, CONTENT_BOTTOM - y2, actors, t("count"))
        self.finish(slide, n, [f"{h.get('object_id')}" for h in human[-6:]], t("section_6"))

    def slide_dataflow(self, n: int) -> None:
        ctx, t, lang = self.ctx, self.t, self.lang
        df = ctx["dataflow"]
        summ = df["summary"]
        slide = self.new_slide(t("px_dataflow"), t("section_8"))
        gap = 0.3
        cw = (SLIDE_W - 2 * MARGIN - 2 * gap) / 3
        ch = 3.05
        left_nothing = not summ.get("n_external")
        cards = [
            (t("px_stayed_local"), ACCENT, [t("s8_stays_items")]),
            (t("px_left"), PASS if left_nothing else "7B1FA2", [t("px_nothing_left")] if left_nothing else [t("px_left_text", n_external=summ.get("n_external", 0), bytes=_thousands(summ.get("bytes_external", 0), lang), model=df.get("external_model"), n_blocked=summ.get("n_blocked", 0)), t("s8_leaves_items")]),
            (t("px_model_used"), MUTED, [f"{t('s8_local_model')}: {df.get('local_provider')} / {df.get('local_model')} — {_thousands(summ.get('n_local', 0), lang)} {t('calls_local').lower()}", f"{t('s8_external_model')}: {df.get('external_provider')} / {df.get('external_model')} — {_thousands(summ.get('n_external', 0), lang)} {t('calls_external').lower()}", f"{t('s8_profile')}: {df.get('profile')}"]),
        ]
        for k, (head, color, body) in enumerate(cards):
            x = MARGIN + k * (cw + gap)
            self.rect(slide, x, CONTENT_TOP, cw, ch, fill=SOFT)
            self.rect(slide, x, CONTENT_TOP, cw, 0.08, fill=color)
            self.text(slide, x + 0.1, CONTENT_TOP + 0.15, cw - 0.2, 0.42, [{"text": head, "size": 14, "bold": True, "color": color if color != MUTED else INK, "after": 0}], anchor="middle", role="card-head")
            self.text(slide, x + 0.1, CONTENT_TOP + 0.6, cw - 0.2, ch - 0.7, [{"text": b, "size": 11.5, "after": 5} for b in body], role="card-body")
        y = CONTENT_TOP + ch + 0.25
        rows = [[_thousands(summ.get("n_local", 0), lang), _thousands(summ.get("n_external", 0), lang), _thousands(summ.get("n_blocked", 0), lang), _thousands(summ.get("n_fallback", 0), lang), _thousands(summ.get("bytes_external", 0), lang)]]
        self.table(slide, MARGIN, y, SLIDE_W - 2 * MARGIN, [t("calls_local"), t("calls_external"), t("calls_blocked"), t("calls_fallback"), t("bytes_external")], rows, [20, 20, 20, 20, 20], size=13, row_h=0.45)
        y += 0.36 + 0.45 + 0.12
        foot = []  # round 6: who wrote the explanations, and the guard demonstration on this run (when made)
        if df.get("coverage"):
            foot.append({"text": str(df["coverage"].get("sentence") or ""), "size": 10.5, "color": INK, "after": 3})
        gd = df.get("guard_demo")
        if gd:
            first = next(iter(gd.get("unsafe") or []), "")
            foot.append({"text": f"{gd.get('title')}: {first}".strip(": "), "size": 10, "color": MUTED, "after": 0})
        if not foot:
            foot = [{"text": t("s8_intro"), "size": 10.5, "color": MUTED, "after": 0}]
        self.text(slide, MARGIN - INSET, y, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - y + 0.05, foot, role="body")
        self.finish(slide, n, [r.get("id") for r in (df.get("ledger") or [])[:12]], t("section_8"))

    def slide_assessor(self, n: int) -> None:
        ctx, t = self.ctx, self.t
        a = ctx.get("assessor") or {}
        slide = self.new_slide(t("px_assessor"), t("section_assessor"))
        pts = a.get("curve_points") or []
        lw = 7.0 if len(pts) >= 2 else SLIDE_W - 2 * MARGIN
        y = CONTENT_TOP
        verdicts = a.get("verdicts") or []
        ev: list[str] = []
        if not a:
            self.text(slide, MARGIN - INSET, y, lw, 0.8, [{"text": t("assessor_no"), "size": 12, "color": MUTED}], role="body")
        for v in verdicts[:2]:
            color = PASS if v.get("state") == "pass" else (WARN if v.get("state") == "warn" else MUTED)
            self.rect(slide, MARGIN, y, lw, 1.5, fill=SOFT)
            self.rect(slide, MARGIN, y, 0.08, 1.5, fill=color)
            self.text(slide, MARGIN + 0.15, y + 0.05, lw - 1.6, 0.42, [{"text": v.get("question"), "size": 13, "bold": True, "after": 0}], anchor="middle", role="card-head")
            self.chip(slide, MARGIN + lw - 1.4, y + 0.1, str(v.get("answer")), color, w=1.25)
            self.text(slide, MARGIN + 0.15, y + 0.5, lw - 0.3, 0.98, [{"text": v.get("why") or "", "size": 11, "after": 0}], role="card-body")
            y += 1.65
        rest = [{"text": a.get("summary"), "size": 11.5, "after": 5}] if a.get("summary") else []
        for r in (a.get("recommendations") or [])[:2]:
            rest.append({"text": f"{t('recommendations')}: {r.get('action')}" + (f" — {r.get('text')}" if r.get("text") else ""), "size": 11, "bullet": "•", "after": 3})
            ev += ids_in(r.get("evidence"), r.get("id"))
        if rest and y < CONTENT_BOTTOM - 0.5:
            self.text(slide, MARGIN - INSET, y, lw + INSET, CONTENT_BOTTOM - y, rest, role="body")
        if len(pts) >= 2:
            x = MARGIN + lw + 0.45
            w = SLIDE_W - MARGIN - x
            y2 = self.heading(slide, x, CONTENT_TOP - 0.05, w, t("learning_curve"))
            self.xy_chart(slide, x - 0.1, y2, w + 0.1, CONTENT_BOTTOM - y2, pts, t("learning_curve"))
        self.finish(slide, n, ev, t("section_assessor"))

    def _eval_metrics(self) -> list[tuple[str, str]]:
        ev = self.ctx.get("evaluation") or {}
        out: list[tuple[str, str]] = []

        def add(name: str, v: Any) -> None:
            if isinstance(v, (int, float)) and not isinstance(v, bool) and len(out) < 10:
                whole = float(v).is_integer() and abs(v) >= 1000
                out.append((name.replace("_", " "), _thousands(int(v), self.lang) if whole else _num(v, 3)))

        for k, v in ev.get("scalars") or []:
            add(str(k), v)
        for tb in ev.get("tables") or []:
            for key, vals in (tb.get("rows") or [])[:2]:
                for col, v in zip(tb.get("cols") or [], vals):
                    if isinstance(v, dict):
                        for kk, vv in v.items():
                            add(f"{key} · {kk}", vv)
                    else:
                        add(f"{key} · {col}", v)
        return out

    def slide_evaluation(self, n: int) -> None:
        t = self.t
        slide = self.new_slide(t("px_evaluation"), t("section_eval"))
        metrics = self._eval_metrics()
        lw = 7.2
        if metrics:
            self.table(slide, MARGIN, CONTENT_TOP + 0.05, lw, [t("key"), t("value")], [[k, v] for k, v in metrics], [70, 30], size=11.5, row_h=0.4)
        x = MARGIN + lw + 0.5
        w = SLIDE_W - MARGIN - x
        notes = [v for k, v in (self.ctx["evaluation"].get("scalars") or []) if isinstance(v, str) and len(v) > 20][:2]
        self.text(slide, x - INSET, CONTENT_TOP, w + INSET, CONTENT_BOTTOM - CONTENT_TOP, [{"text": t("eval_intro"), "size": 12.5, "after": 8}] + [{"text": s, "size": 11, "color": "374151", "bullet": "•", "after": 4} for s in notes], role="body")
        self.finish(slide, n, [], t("section_eval"))

    def slide_next(self, n: int) -> None:
        ctx, t = self.ctx, self.t
        slide = self.new_slide(t("px_next"), t("llm_uncertainty"))
        half = (SLIDE_W - 2 * MARGIN - 0.6) / 2
        y = self.heading(slide, MARGIN, CONTENT_TOP - 0.05, half, t("px_open_points"))
        pts = open_points(ctx, 5)
        self.text(slide, MARGIN - INSET, y, half + INSET, CONTENT_BOTTOM - y, [{"text": p, "size": 12, "bullet": "•", "after": 7} for p in pts], role="body")
        x = MARGIN + half + 0.6
        y = self.heading(slide, x, CONTENT_TOP - 0.05, half, t("px_next_steps"))
        self.text(slide, x - INSET, y, half + INSET, CONTENT_BOTTOM - y, [{"text": p, "size": 12, "bullet": "1", "after": 7} for p in next_steps(ctx, 5)], role="body")
        self.finish(slide, n, ids_in(pts), t("section_overview"))

    # ------------------------------------------------------------------ build
    def build(self) -> bytes:
        ctx = self.ctx
        self.slide_title()
        n = 2
        plan: list[Any] = [self.slide_summary, self.slide_analysed, self.slide_sensors, self.slide_trust, self.slide_monitoring]
        su = ctx.get("suspicious")
        if su and (su["rows"] or su["point_flags"]):
            plan.append(self.slide_suspicious)
        diags = top_diagnoses(ctx, 3)
        for i, d in enumerate(diags, start=1):
            plan.append(lambda k, d=d, i=i: self.slide_diagnosis(k, d, i, len(diags)))
        if ctx["detect"].get("patterns"):
            plan.append(self.slide_patterns)
        plan += [self.slide_human, self.slide_dataflow, self.slide_assessor]
        if ctx.get("evaluation"):
            plan.append(self.slide_evaluation)
        plan.append(self.slide_next)
        for fn in plan:
            fn(n)
            n += 1
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


# ----------------------------------------------------------------------------- public API
def render_pptx(context: dict[str, Any]) -> bytes:
    """The .pptx bytes for a report context (collect() / export_context())."""
    return _Deck(context).build()


def generate_pptx(ws: Workspace, settings: Optional[Settings] = None, lang: str = "en", out_path: Optional[str | Path] = None, context: Optional[dict[str, Any]] = None, log: bool = True) -> Path:
    """Write the deck (default <run>/report_<lang>.pptx) and return its path. The language model is never called: a
    model-written summary is used only when one is already stored for this run and language."""
    settings = settings or ws.settings or get_settings()
    lang = normalize_lang(lang)
    ctx = context if context is not None and context.get("lang") == lang else export_context(ws, settings, lang)
    out = Path(out_path) if out_path else export_path(ws, lang, "pptx")
    deck = _Deck(ctx)
    data = deck.build()
    write_atomic_bytes(out, data)
    if log:
        log_export(ws, out, lang, "pptx", {"slides": len(deck.prs.slides), **deck.stats})
    return out


__all__ = ["generate_pptx", "render_pptx", "text_budget", "cell_capacity", "chars_per_line", "fit_paragraphs", "ELLIPSIS"]
